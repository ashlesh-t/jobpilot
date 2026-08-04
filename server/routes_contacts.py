"""HR/recruiter contacts — CRUD plus a CSV/XLSX bulk import.

Matched to jobs by company name (core/repo/_company_match.py) so a job's detail view
can show "you know someone here" without the user doing anything per-job.
"""
from __future__ import annotations

import csv
import io
import sys
import tempfile
from pathlib import Path

from fastapi import APIRouter, File, HTTPException, UploadFile
from pydantic import BaseModel

REPO_DIR = Path(__file__).resolve().parent.parent
if str(REPO_DIR) not in sys.path:
    sys.path.insert(0, str(REPO_DIR))

from core.repo import contacts as contacts_repo  # noqa: E402

router = APIRouter(prefix="/api/contacts", tags=["contacts"])

MAX_BYTES = 10 * 1024 * 1024  # 10 MB — a 300-row contact sheet is a few hundred KB
ALLOWED_SUFFIXES = {".csv", ".xlsx"}


class ContactCreate(BaseModel):
    company: str
    name: str = ""
    email: str = ""
    role: str = ""


class ContactUpdate(BaseModel):
    company: str | None = None
    name: str | None = None
    email: str | None = None
    role: str | None = None


def _parse_csv(raw: bytes) -> list[dict]:
    text = raw.decode("utf-8-sig", errors="replace")
    return list(csv.DictReader(io.StringIO(text)))


def _parse_xlsx(path: Path) -> list[dict]:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    try:
        header = [str(h or "").strip() for h in next(rows)]
    except StopIteration:
        return []
    out = []
    for row in rows:
        if not any(row):
            continue
        out.append({header[i]: (row[i] if i < len(row) else "") for i in range(len(header))})
    return out


@router.get("")
async def list_contacts(company: str | None = None):
    return {"contacts": contacts_repo.list_all(company=company)}


@router.post("")
async def create_contact(req: ContactCreate):
    return contacts_repo.create(company=req.company, name=req.name, email=req.email,
                               role=req.role, source="manual")


@router.patch("/{contact_id}")
async def update_contact(contact_id: int, req: ContactUpdate):
    updated = contacts_repo.update(contact_id, **req.model_dump(exclude_none=True))
    if updated is None:
        raise HTTPException(status_code=404, detail="not found")
    return updated


@router.delete("/{contact_id}")
async def delete_contact(contact_id: int):
    if not contacts_repo.remove(contact_id):
        raise HTTPException(status_code=404, detail="not found")
    return {"ok": True}


@router.post("/import")
async def import_contacts(file: UploadFile = File(...)):
    suffix = Path(file.filename or "").suffix.lower()
    if suffix not in ALLOWED_SUFFIXES:
        raise HTTPException(
            status_code=400,
            detail=f"{suffix or 'that file type'} isn't supported — use .csv or .xlsx")

    with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
        size = 0
        while chunk := await file.read(1024 * 256):
            size += len(chunk)
            if size > MAX_BYTES:
                tmp.close()
                Path(tmp.name).unlink(missing_ok=True)
                raise HTTPException(status_code=413, detail="that file is larger than 10 MB")
            tmp.write(chunk)
        temp_path = Path(tmp.name)

    try:
        if suffix == ".csv":
            rows = _parse_csv(temp_path.read_bytes())
        else:
            rows = _parse_xlsx(temp_path)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=400, detail=f"could not read that file: {exc}")
    finally:
        temp_path.unlink(missing_ok=True)

    if not rows:
        raise HTTPException(status_code=400, detail="no rows found in that file")

    result = contacts_repo.bulk_import(rows)
    result["contacts"] = contacts_repo.list_all()
    return result
